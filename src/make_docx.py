"""Convert REPORT.md to REPORT.docx with the figures embedded.

1. Copy the markdown, drop the hand-written contents list (Word gets a real table of
   contents) and turn each "Figures: `a.png`, `b.png`." line into embedded images.
2. Convert with pandoc.
3. Tidy the Word file with python-docx: fonts, table borders and header shading, image
   widths and margins.

    python src/make_docx.py
"""
import re
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
FIG = ROOT / "results" / "figures"
CAPTIONS = {
    "price_accuracy.png": "Next-day price error and share of forecasts within 2%, by model",
    "price_paths_AAPL.png": "AAPL one-day-ahead forecasts vs actual closing prices",
    "shap_global.png": "Top 15 features by mean |SHAP|",
    "shap_dependence.png": "How the top features move the forecast (SHAP dependence)",
    "local_explanations.png": "Local SHAP explanations in dollars for individual stock-days",
    "faithfulness.png": "Faithfulness: forecast change when top-SHAP vs random features are removed",
    "stability_sanity.png": "Stability across seeds and retraining, and comparison with no-signal models",
    "regime_family_shares.png": "Share of SHAP attribution by feature family across market regimes",
    "sector_family_shares.png": "Share of SHAP attribution by feature family and GICS sector",
    "lime_vs_shap_local.png": "SHAP vs LIME explanations for the same forecasts",
    "lime_reliability.png": "LIME vs SHAP: run-to-run stability, faithfulness and local fit",
    "explainer_agreement.png": "Feature importance ranks from TreeSHAP, LIME, permutation importance and MDI",
}
ACCENT = RGBColor(0x1F, 0x3A, 0x5F)


LIST_ITEM = re.compile(r"^\s*(\*|-|\d+\.)\s")


def blank_line_before_lists(text):
    """Pandoc only starts a list after a blank line; the markdown often has a list straight
    after a bold lead-in line, which would otherwise render as run-on text with literal '*'."""
    out, prev = [], ""
    for line in text.split("\n"):
        if LIST_ITEM.match(line) and prev.strip() and not LIST_ITEM.match(prev) and not prev.startswith((" ", "\t", "|", ">")):
            out.append("")
        out.append(line)
        prev = line
    return "\n".join(out)


def prepare_markdown(text):
    text = re.sub(r"## Contents\n.*?\n---\n\n", "", text, count=1, flags=re.S)
    text = blank_line_before_lists(text)
    n = 0

    def figures(match):
        nonlocal n
        blocks = []
        for name in re.findall(r"`([^`]+\.png)`", match.group(0)):
            n += 1
            blocks.append(f"![Figure {n}. {CAPTIONS.get(name, name)}]({FIG / name})")
        return "\n\n".join(blocks)

    return re.sub(r"^Figures: .*$", figures, text, flags=re.M)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_table_borders(table, color="A6B1BF"):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tbl_pr.append(borders)


def set_column_widths(table, total):
    """Explicit widths (pandoc's can collapse to zero): proportional to the longest text in
    each column, with a floor so short columns stay readable."""
    n_cols = len(table.columns)
    lengths = [max(len(table.cell(r, c).text) for r in range(len(table.rows))) for c in range(n_cols)]
    weights = [min(max(l, 16), 60) for l in lengths]
    widths = [int(total * w / sum(weights)) for w in weights]
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    for tag in ("w:tblW", "w:tblLayout"):
        old = tbl_pr.find(qn(tag))
        if old is not None:
            tbl_pr.remove(old)
    tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:w"), str(int(total.twips if hasattr(total, "twips") else total / 635)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_w)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    grid = tbl.find(qn("w:tblGrid"))
    if grid is not None:
        for col, w in zip(grid.findall(qn("w:gridCol")), widths):
            col.set(qn("w:w"), str(int(w / 635)))
    for row in table.rows:
        for c, cell in enumerate(row.cells[:n_cols]):
            cell.width = widths[c]


def polish(path):
    doc = Document(path)
    for section in doc.sections:
        section.page_width, section.page_height = Cm(21.0), Cm(29.7)
        section.left_margin = section.right_margin = Cm(2.0)
        section.top_margin = section.bottom_margin = Cm(2.0)

    styles = doc.styles
    for name in ("Normal", "Body Text", "First Paragraph", "Compact"):
        if name in styles:
            styles[name].font.name = "Calibri"
            styles[name].font.size = Pt(10.5)
    for level, size in ((1, 20), (2, 15), (3, 12.5)):
        name = f"Heading {level}"
        if name in styles:
            st = styles[name]
            st.font.name = "Calibri"
            st.font.size = Pt(size)
            st.font.bold = True
            st.font.color.rgb = ACCENT
    for name in ("Title", "Subtitle"):
        if name in styles:
            styles[name].font.name = "Calibri"
            styles[name].font.color.rgb = ACCENT

    usable = Cm(17.0)
    for shape in doc.inline_shapes:
        ratio = shape.height / shape.width
        shape.width = usable
        shape.height = int(usable * ratio)

    # Word should refresh the table of contents when the file is opened.
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    doc.settings.element.append(update)

    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(table)
        set_column_widths(table, usable)
        for r, row in enumerate(table.rows):
            for cell in row.cells:
                if r == 0:
                    set_cell_shading(cell, "DCE6F1")
                for para in cell.paragraphs:
                    para.paragraph_format.space_after = Pt(2)
                    for run in para.runs:
                        run.font.size = Pt(9)
                        run.font.name = "Calibri"
                        if r == 0:
                            run.font.bold = True
    doc.save(path)


def main():
    md = prepare_markdown((ROOT / "REPORT.md").read_text())
    # Title and subtitle become document metadata so they appear above the table of contents;
    # section headings then move up one level (## -> Heading 1).
    title = re.search(r"^# (.+)$", md, flags=re.M).group(1)
    subtitle = re.search(r"^\*([^*\n]+)\*$", md, flags=re.M).group(1)
    md = re.sub(r"^# .+\n+\*[^*\n]+\*\n+---\n", "", md, count=1, flags=re.M)
    out = ROOT / "REPORT.docx"
    with tempfile.NamedTemporaryFile("w", suffix=".md", delete=False) as tmp:
        tmp.write(md)
    subprocess.run(["pandoc", tmp.name, "-f", "markdown", "-t", "docx", "--toc", "--toc-depth=2",
                    "--shift-heading-level-by=-1", "-M", f"title={title}", "-M", f"subtitle={subtitle}",
                    "-M", "toc-title=Contents", "-o", str(out)], check=True)
    polish(out)
    print("wrote", out)


if __name__ == "__main__":
    main()
