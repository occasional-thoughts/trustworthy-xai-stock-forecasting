"""Rendering primitives for the VIT project report.

Typography and page setup follow the School of Computer Science and Engineering
template: A4, Times New Roman 12, line spacing 1.15 in the body, figure captions below
figures and table captions above tables at 10 pt. Front matter is numbered in lower-case
roman and the chapters restart at 1.

Used by make_vit_report.py; not run on its own.
"""
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BODY = "Times New Roman"
MONO = "Consolas"
BODY_PT = 12
CAPTION_PT = 10
LINE = 1.15
HEAD_FILL = "D9E2F3"
RULE = "7F7F7F"


# ---------------------------------------------------------------- low-level helpers

def _field(paragraph, instr):
    """Insert a Word field (used for page numbers, which Word fills in itself)."""
    run = paragraph.add_run()
    for tag, attr, text in (("w:fldChar", "w:fldCharType", "begin"),
                            ("w:instrText", "xml:space", "preserve"),
                            ("w:fldChar", "w:fldCharType", "end")):
        el = OxmlElement(tag)
        el.set(qn(attr), text)
        if tag == "w:instrText":
            el.text = instr
        run._r.append(el)
    return run


def _page_numbering(section, fmt, start=None):
    sect_pr = section._sectPr
    old = sect_pr.find(qn("w:pgNumType"))
    if old is not None:
        sect_pr.remove(old)
    el = OxmlElement("w:pgNumType")
    el.set(qn("w:fmt"), fmt)
    if start is not None:
        el.set(qn("w:start"), str(start))
    sect_pr.append(el)


def _shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _borders(table, color=RULE):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tbl_pr.append(borders)


def _row_rules(table):
    """Repeat the header row on every page and stop rows splitting across a page break."""
    for i, row in enumerate(table.rows):
        tr_pr = row._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        tr_pr.append(cant)
        if i == 0:
            hdr = OxmlElement("w:tblHeader")
            tr_pr.append(hdr)


def _widths(table, total, weights=None):
    """Explicit fixed widths; pandoc-free tables otherwise collapse unevenly."""
    n = len(table.columns)
    if weights is None:
        lengths = [max(len(table.cell(r, c).text) for r in range(len(table.rows))) for c in range(n)]
        weights = [min(max(l, 12), 54) for l in lengths]
    widths = [int(total * w / sum(weights)) for w in weights]
    tbl_pr = table._tbl.tblPr
    for tag in ("w:tblW", "w:tblLayout"):
        old = tbl_pr.find(qn(tag))
        if old is not None:
            tbl_pr.remove(old)
    tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:w"), str(int(total / 635)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_w)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    grid = table._tbl.find(qn("w:tblGrid"))
    if grid is not None:
        for col, w in zip(grid.findall(qn("w:gridCol")), widths):
            col.set(qn("w:w"), str(int(w / 635)))
    for row in table.rows:
        for c, cell in enumerate(row.cells[:n]):
            cell.width = widths[c]


# ---------------------------------------------------------------- document skeleton

def setup(doc):
    """Page size, margins and the base character style."""
    for section in doc.sections:
        section.page_width, section.page_height = Cm(21.0), Cm(29.7)
        section.left_margin, section.right_margin = Cm(2.96), Cm(2.33)
        section.top_margin, section.bottom_margin = Cm(2.43), Cm(2.0)
    normal = doc.styles["Normal"]
    normal.font.name = BODY
    normal.font.size = Pt(BODY_PT)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    rpr = normal.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attr), BODY)
    pf = normal.paragraph_format
    pf.line_spacing = LINE
    pf.space_after = Pt(6)
    for name in ("List Bullet", "List Number"):
        if name in doc.styles:
            st = doc.styles[name]
            st.font.name = BODY
            st.font.size = Pt(BODY_PT)
            st.paragraph_format.line_spacing = LINE
            st.paragraph_format.space_after = Pt(3)


def footer_page_numbers(section):
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = _field(p, "PAGE")
    run.font.name = BODY
    run.font.size = Pt(11)
    for r in p.runs:
        r.font.name = BODY
        r.font.size = Pt(11)


def body_section(doc):
    """Start the numbered chapters: a new section whose page numbers restart at 1."""
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    section.left_margin, section.right_margin = Cm(2.96), Cm(2.33)
    section.top_margin, section.bottom_margin = Cm(2.43), Cm(2.0)
    _page_numbering(section, "decimal", start=1)
    footer_page_numbers(section)
    return section


def front_matter_numbering(doc):
    _page_numbering(doc.sections[0], "lowerRoman", start=1)
    footer_page_numbers(doc.sections[0])


# ---------------------------------------------------------------- content primitives

def para(doc, text="", align="justify", size=BODY_PT, bold=False, italic=False, space_after=6,
         line=LINE, font=BODY, indent=None, keep_with_next=False):
    p = doc.add_paragraph()
    p.alignment = {"justify": WD_ALIGN_PARAGRAPH.JUSTIFY, "center": WD_ALIGN_PARAGRAPH.CENTER,
                   "left": WD_ALIGN_PARAGRAPH.LEFT, "right": WD_ALIGN_PARAGRAPH.RIGHT}[align]
    pf = p.paragraph_format
    pf.line_spacing = line
    pf.space_after = Pt(space_after)
    pf.keep_with_next = keep_with_next
    if indent is not None:
        pf.left_indent = Cm(indent)
    if text:
        run = p.add_run(text)
        run.font.name = font
        run.font.size = Pt(size)
        run.bold = bold
        run.italic = italic
    return p


def bullets(doc, items, numbered=False, size=BODY_PT):
    for i, item in enumerate(items, start=1):
        if numbered:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1.0)
            p.paragraph_format.first_line_indent = Cm(-1.0)
            text = f"{i}.\t{item}"
        else:
            p = doc.add_paragraph(style="List Bullet")
            text = item
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = LINE
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(text)
        run.font.name = BODY
        run.font.size = Pt(size)


def code(doc, text, size=8):
    for line in text.split("\n"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf = p.paragraph_format
        pf.line_spacing = 1.0
        pf.space_after = Pt(0)
        pf.left_indent = Cm(0.8)
        run = p.add_run(line if line else " ")
        run.font.name = MONO
        run.font.size = Pt(size)
    para(doc, "", space_after=4)


def caption(doc, text, size=CAPTION_PT, keep_with_next=False):
    p = para(doc, text, align="center", size=size, space_after=10, line=1.0,
             keep_with_next=keep_with_next)
    return p


def table(doc, header, rows, widths=None, size=9.5, align_left=(0,)):
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    usable = Cm(15.71)
    for j, text in enumerate(header):
        cell = t.cell(0, j)
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.0
        run = p.add_run(str(text))
        run.font.name = BODY
        run.font.size = Pt(size)
        run.bold = True
        _shade(cell, HEAD_FILL)
    for i, row in enumerate(rows, start=1):
        for j, text in enumerate(row):
            cell = t.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = (WD_ALIGN_PARAGRAPH.LEFT if j in align_left else WD_ALIGN_PARAGRAPH.CENTER)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.0
            run = p.add_run(str(text))
            run.font.name = BODY
            run.font.size = Pt(size)
    _borders(t)
    _row_rules(t)
    _widths(t, usable.emu if hasattr(usable, "emu") else int(usable), widths)
    para(doc, "", space_after=4)
    return t


def figure(doc, path, width_cm=14.0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    return p
